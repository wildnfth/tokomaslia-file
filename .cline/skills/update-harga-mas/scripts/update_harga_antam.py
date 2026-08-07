# -*- coding: utf-8 -*-
"""
SKILL: update-harga-antam (v2 fleksibel)
========================================
Mengelola blok harga ANTAM di sheet ANTAM, file TEMPLATE HARGA MAS2.xlsx.

DUA MODE:
  1) add   (default) - TAMBAH blok harian baru: menyalin blok terakhir ke
                       bawah (pola 1 baris kosong), menyesuaikan harga sesuai
                       target/gramasi yg diminta, mengganti label tanggal.
  2) update - MENGUBAH harga pada blok TERAKHIR yg SUDAH ADA (di tempat,
              TANPA menambah blok baru). Berguna utk menyesuaikan harga
              kolom tertentu tanpa blok harian baru.

TARGET SELEKTIF (--targets):
  antam      = semua kolom ANTAM (RETRO, RANDOM, 2025, 2026)  [default]
  retro      = kolom B (RETRO)
  random     = kolom C (RANDOM)
  2025       = kolom D
  2026       = kolom E
  galeri24   = kolom H
  ubs        = kolom L (UBS Batik)
  Bisa gabung dgn koma, mis. --targets "2025,ubs".

FILTER GRAMASI (--grams):
  --grams "1,2,5"  hanya ubah baris dgn gramasi tsb.

CONTOH:
  python update_harga_antam.py <file> --date "8 AGUSTUS 2026" --step 50
  python update_harga_antam.py <file> --date "8 AGUSTUS 2026" --step 30 --targets ubs
  python update_harga_antam.py <file> --date "8 AGUSTUS 2026" --step 40 --targets 2025 --grams "1,2"
  python update_harga_antam.py <file> --mode update --step 25 --targets galeri24
  # selalu mulai dgn --dry-run
"""
import argparse
import copy
import os
import shutil
import sys
from datetime import datetime

from openpyxl import load_workbook

HEADER_MARK = "HARGA ANTAM"
# Kolom harga -> kolom gram masing-masing
GRAM_COL = {2: 1, 3: 1, 4: 1, 5: 1, 8: 7, 12: 11}
# Nama target -> daftar kolom harga
TARGET_MAP = {
    "retro": [2],
    "random": [3],
    "2025": [4],
    "2026": [5],
    "galeri24": [8],
    "ubs": [12],
}
ANTAM_NAMES = ["retro", "random", "2025", "2026"]


def find_header_rows(ws):
    """Baris-baris yg memuat header blok 'HARGA ANTAM ...' di kolom A."""
    rows = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip().startswith(HEADER_MARK):
            rows.append(r)
    return rows


def copy_style(src, dst):
    """Salin style (font/fill/border/alignment/number_format/dll) persis."""
    if src.has_style:
        dst._style = copy.copy(src._style)


def resolve_columns(targets_spec):
    """Kembalikan (daftar kolom harga, set label utk laporan)."""
    cols, labels = [], set()
    if not targets_spec:
        targets_spec = "antam"
    for t in targets_spec.split(","):
        t = t.strip().lower()
        if t == "antam":
            for name in ANTAM_NAMES:
                labels.add(name)
                cols.extend(TARGET_MAP[name])
        elif t in TARGET_MAP:
            labels.add(t)
            cols.extend(TARGET_MAP[t])
        else:
            print("[warn] target tidak dikenal, dilewati:", t)
    cols = list(dict.fromkeys(cols))
    return cols, labels


def resolve_grams(grams_spec):
    """Kembalikan set gramasi yg diizinkan (None = semua)."""
    if not grams_spec:
        return None
    out = set()
    for g in grams_spec.split(","):
        try:
            out.add(float(g.strip()))
        except ValueError:
            print("[warn] gramasi tidak valid, dilewati:", g)
    return out or None

def build_new_value(ws, wsv, sr, c, gram_filter, old_header, new_header,
                    old_date, new_date, step, selected_cols):
    """Hitung nilai sel tujuan utk mode add.

    - Teks (header/label tanggal): ganti label.
    - Kolom harga: bila kolom TERPILIH & gramasi cocok -> +step*gram;
      selain itu -> salin nilai sumber apa adanya (tanpa kenaikan).
    """
    v = ws.cell(sr, c).value
    if isinstance(v, str):
        if old_header and v == old_header:
            return new_header
        if old_date and old_date in v:
            return v.replace(old_date, new_date)
        return v
    if c in GRAM_COL:
        gsrc = ws.cell(sr, GRAM_COL[c]).value
        try:
            g = float(gsrc)
        except (TypeError, ValueError):
            return v
        if c not in selected_cols:
            return v
        if gram_filter is not None and g not in gram_filter:
            return v
        base = v
        if isinstance(v, str) and v.startswith("="):
            base = wsv.cell(sr, c).value
        try:
            base = float(base)
        except (TypeError, ValueError):
            return v
        return base + step * g
    return v


def apply_inplace(ws, wsv, rows, step, selected_cols, gram_filter):
    """Mode update: ubah harga blok terakhir DI TEMPAT, hanya kolom & gramasi
    yg dipilih. Kolom lain dibiarkan utuh."""
    changed = []
    for r in rows:
        for c in selected_cols:
            gsrc = ws.cell(r, GRAM_COL[c]).value
            try:
                g = float(gsrc)
            except (TypeError, ValueError):
                continue
            if gram_filter is not None and g not in gram_filter:
                continue
            v = ws.cell(r, c).value
            base = v
            if isinstance(v, str) and v.startswith("="):
                base = wsv.cell(r, c).value
            try:
                base = float(base)
            except (TypeError, ValueError):
                continue
            nv = base + step * g
            changed.append((r, c, g, base, nv))
            ws.cell(r, c).value = nv
    return changed


def main():
    ap = argparse.ArgumentParser(description="Skill update blok harga ANTAM (fleksibel).")
    ap.add_argument("file", help="Path file .xlsx")
    ap.add_argument("--date", help="Label tanggal baru (wajib utk mode add), mis. '8 AGUSTUS 2026'")
    ap.add_argument("--step", type=int, default=50, help="Kenaikan per gram (boleh negatif), default 50")
    ap.add_argument("--targets", default="antam",
                    help="Kolom yg diubah: antam/retro/random/2025/2026/galeri24/ubs (koma utk banyak)")
    ap.add_argument("--grams", default=None, help="Hanya gramasi tertentu, mis. '1,2,5'")
    ap.add_argument("--mode", choices=["add", "update"], default="add",
                    help="add=tambah blok baru (default), update=ubah blok terakhir di tempat")
    ap.add_argument("--sheet", default="ANTAM", help="Nama sheet (default ANTAM)")
    ap.add_argument("--dry-run", action="store_true", help="Cek & tampilkan rencana saja")
    ap.add_argument("--no-backup", action="store_true", help="Tanpa backup otomatis")
    args = ap.parse_args()

    if not os.path.exists(args.file):
        sys.exit("File tidak ditemukan: %s" % args.file)
    if args.mode == "add" and not args.date:
        sys.exit("Mode add butuh --date. Gunakan --mode update utk ubah blok yg ada.")

    selected_cols, labels = resolve_columns(args.targets)
    gram_filter = resolve_grams(args.grams)
    if not selected_cols:
        sys.exit("Tidak ada target valid.")

    if not args.no_backup and not args.dry_run:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        base, ext = os.path.splitext(args.file)
        shutil.copy2(args.file, "%s_BACKUP_%s%s" % (base, ts, ext))
        print("[backup] dibuat otomatis.")

    wb = load_workbook(args.file, data_only=False)
    wbv = load_workbook(args.file, data_only=True)
    if args.sheet not in wb.sheetnames:
        sys.exit("Sheet '%s' tidak ditemukan. Ada: %s" % (args.sheet, wb.sheetnames))
    ws = wb[args.sheet]
    wsv = wbv[args.sheet]

    headers = find_header_rows(ws)
    if not headers:
        sys.exit("Tidak ada blok 'HARGA ANTAM ...' di sheet %s." % args.sheet)
    src_h = headers[-1]
    spacing = (headers[-1] - headers[-2]) if len(headers) >= 2 else 12
    src_last = ws.max_row
    old_header = str(ws.cell(src_h, 1).value)
    old_date = old_header.split("-", 1)[-1].strip() if "-" in old_header else ""

    data_start = src_h + 2  # baris pertama data (setelah header + sub-header)
    data_rows = list(range(data_start, src_last + 1))
    block_rows = list(range(src_h, src_last + 1))  # seluruh blok termasuk header
    print("[info] mode=%s | target: %s | kolom: %s" % (args.mode, sorted(labels), selected_cols))
    print("[info] step %d/gram | gramasi: %s" % (args.step, gram_filter or "semua"))
    print("[info] blok sumber: %s-%s | header: %s" % (src_h, src_last, old_header))


    if args.mode == "add":
        dst_h = src_h + spacing
        n_rows = len(block_rows)
        new_header = HEADER_MARK + " - " + args.date
        if args.dry_run:
            print("[dry-run/add] blok baru: %s-%s -> %s" % (dst_h, dst_h + n_rows - 1, new_header))
            print("[dry-run] selesai (mode add, belum disimpan).")
            return
        max_col = 13
        for r in block_rows:
            for c in range(1, ws.max_column + 1):
                if ws.cell(r, c).value is not None:
                    max_col = max(max_col, c)
        for m in ws.merged_cells.ranges:
            if m.max_col > max_col and not (m.max_row < src_h or m.min_row > src_last):
                max_col = max(max_col, m.max_col)
        dst_rows = list(range(dst_h, dst_h + n_rows))
        for di, sr in enumerate(block_rows):
            dr = dst_rows[di]
            for c in range(1, max_col + 1):
                s = ws.cell(sr, c)
                d = ws.cell(dr, c)
                d.value = build_new_value(ws, wsv, sr, c, gram_filter, old_header,
                                          new_header, old_date, args.date, args.step,
                                          selected_cols)
                copy_style(s, d)
        for di, sr in enumerate(block_rows):
            dr = dst_rows[di]
            if sr in ws.row_dimensions and ws.row_dimensions[sr].height is not None:
                ws.row_dimensions[dr].height = ws.row_dimensions[sr].height
        for m in list(ws.merged_cells.ranges):
            if not (m.max_row < src_h or m.min_row > src_last):
                ws.merge_cells(start_row=m.min_row + spacing, start_column=m.min_col,
                               end_row=m.max_row + spacing, end_column=m.max_col)
        wb.save(args.file)
        print("[ok] %s -> blok baru %s-%s (%s)" % (args.file, dst_h, dst_h + n_rows - 1, new_header))
        return


    # mode update
    rows = []
    for r in data_rows:
        rows.append(r)
        if isinstance(ws.cell(r, 7).value, str) and ws.cell(r, 7).value.strip() == "MERAH = KOSONG":
            break
    if args.dry_run:
        print("[dry-run/update] baris yg akan diubah (kolom %s):" % selected_cols)
        for r in rows:
            for c in selected_cols:
                g = ws.cell(r, GRAM_COL[c]).value
                try:
                    g = float(g)
                except (TypeError, ValueError):
                    continue
                if gram_filter is not None and g not in gram_filter:
                    continue
                v = ws.cell(r, c).value
                base = v
                if isinstance(v, str) and v.startswith("="):
                    base = wsv.cell(r, c).value
                try:
                    base = float(base)
                except (TypeError, ValueError):
                    continue
                print("  row %d | gram %s | col %d | %s -> %s" % (r, g, c, base, base + args.step * g))
        print("[dry-run] selesai (mode update, belum disimpan).")
        return
    changed = apply_inplace(ws, wsv, rows, args.step, selected_cols, gram_filter)
    wb.save(args.file)
    print("[ok] %s -> blok terakhir diupdate: %d sel (target %s)" % (args.file, len(changed), sorted(labels)))


if __name__ == "__main__":
    main()


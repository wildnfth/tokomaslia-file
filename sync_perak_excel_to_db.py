"""
Sync harga perak dari TEMPLATE HARGA PERAK.xlsx ke harga_perhiasan.db (tabel harga_perak).
Jalankan setelah update Excel. Tutup file Excel dulu.
"""
import re
import sqlite3

import openpyxl

DB_PATH = r"D:/TOKO MAS LIA/DOKUMEN/harga_perhiasan.db"
XLSX_PATH = r"D:/TOKO MAS LIA/DOKUMEN/TEMPLATE HARGA PERAK.xlsx"

MONTHS = {
    "JANUARI": 1, "FEBRUARI": 2, "MARET": 3, "APRIL": 4, "MEI": 5, "JUNI": 6,
    "JULI": 7, "AGUSTUS": 8, "SEPTEMBER": 9, "OKTOBER": 10, "NOVEMBER": 11, "DESEMBER": 12,
}


def parse_tanggal_header(text: str) -> str:
    m = re.search(r"TANGGAL\s+(\d{1,2})\s+(\w+)\s+(\d{4})", str(text).upper())
    if not m:
        raise ValueError(f"Tanggal tidak dikenali: {text!r}")
    day, month_name, year = int(m.group(1)), m.group(2), int(m.group(3))
    month = MONTHS[month_name]
    return f"{year:04d}-{month:02d}-{day:02d}"


def parse_berat(text: str) -> float:
    t = str(text).strip().lower()
    m = re.match(r"^(\d+(?:[.,]\d+)?)\s*gr$", t)
    if m:
        return float(m.group(1).replace(",", "."))
    m = re.match(r"^(\d+(?:[.,]\d+)?)\s*kg$", t)
    if m:
        return float(m.group(1).replace(",", ".")) * 1000
    raise ValueError(f"Berat tidak dikenali: {text!r}")


def parse_harga_cell(value):
    if value is None:
        return None, None, None
    if isinstance(value, (int, float)):
        return float(value), None, None
    s = str(value).strip()
    if s in ("-", "—"):
        return None, None, "-"
    m = re.match(r"^([\d,]+)\s*rb/gr\s*$", s, re.I)
    if m:
        return float(m.group(1).replace(",", ".")), None, None
    m = re.match(r"^(\d+(?:[.,]\d+)?)\s*jt\s*$", s, re.I)
    if m:
        return None, int(float(m.group(1).replace(",", ".")) * 1_000_000), None
    if s.lower() == "tanya":
        return None, None, "Tanya"
    raise ValueError(f"Harga tidak dikenali: {value!r}")


def import_section(ws, start_row: int, end_row: int, tanggal_iso: str) -> list:
    records = []
    merek_left = merek_right = None
    for r in range(start_row, end_row + 1):
        a, d = ws.cell(r, 1).value, ws.cell(r, 4).value
        b, e = ws.cell(r, 2).value, ws.cell(r, 5).value
        g, j = ws.cell(r, 7).value, ws.cell(r, 10).value

        if a in ("STAR SILVER", "LOTUS", "ANTAM", "LM PERAK"):
            merek_left = a
        if d in ("MT", "SIMBA", "EURO"):
            merek_right = d

        if a and str(a).strip().endswith(("gr", "kg")) and merek_left:
            berat = parse_berat(a)
            if b is not None and str(b).strip():
                hpg, ht, ket = parse_harga_cell(b)
            elif isinstance(g, (int, float)):
                hpg, ht, ket = float(g), None, None
            else:
                hpg, ht, ket = None, None, None
            if merek_left != "LM PERAK" or hpg or ht or ket:
                records.append((tanggal_iso, merek_left, berat, hpg, ht, ket))

        if d and str(d).strip().endswith(("gr", "kg")) and merek_right:
            berat = parse_berat(d)
            if e is not None and str(e).strip() and str(e).strip() != "-":
                hpg, ht, ket = parse_harga_cell(e)
            elif isinstance(j, (int, float)):
                hpg, ht, ket = float(j), None, None
            elif e == "-":
                hpg, ht, ket = None, None, "-"
            else:
                hpg, ht, ket = None, None, None
            records.append((tanggal_iso, merek_right, berat, hpg, ht, ket))
    return records


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb.active
    sections = []
    for r in range(1, ws.max_row + 1):
        val = ws.cell(r, 1).value
        if val and str(val).upper().startswith("TANGGAL "):
            sections.append((r, parse_tanggal_header(val)))

    conn = sqlite3.connect(DB_PATH)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS harga_perak (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            tanggal TEXT NOT NULL,
            merek TEXT NOT NULL,
            berat_gram REAL NOT NULL,
            harga_per_gram REAL,
            harga_total INTEGER,
            keterangan TEXT,
            UNIQUE(tanggal, merek, berat_gram)
        )""")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_perak_tanggal ON harga_perak(tanggal)")

    for i, (start_row, iso) in enumerate(sections):
        end_row = sections[i + 1][0] - 1 if i + 1 < len(sections) else ws.max_row
        recs = import_section(ws, start_row + 1, end_row, iso)
        conn.execute("DELETE FROM harga_perak WHERE tanggal = ?", (iso,))
        conn.executemany(
            """INSERT INTO harga_perak
               (tanggal, merek, berat_gram, harga_per_gram, harga_total, keterangan)
               VALUES (?,?,?,?,?,?)""",
            recs,
        )
        print(f"{iso}: {len(recs)} baris")
    conn.commit()
    conn.close()
    print("Selesai →", DB_PATH)


if __name__ == "__main__":
    main()
"""
Sync HARGA UBS BATIK dari sheet ANTAM (TEMPLATE HARGA MAS2.xlsx) → harga_ubs_batik.
Jalankan setelah update Excel LM yang punya blok UBS. Tutup file Excel dulu.
"""
from __future__ import annotations

import re
import sqlite3

import openpyxl

DB_PATH = r"D:/TOKO MAS LIA/DOKUMEN/harga_perhiasan.db"
XLSX_PATH = r"D:/TOKO MAS LIA/DOKUMEN/TEMPLATE HARGA MAS2.xlsx"
SHEET = "ANTAM"

MONTHS = {
    "JANUARI": 1,
    "FEBRUARI": 2,
    "MARET": 3,
    "APRIL": 4,
    "MEI": 5,
    "JUNI": 6,
    "JULI": 7,
    "AGUSTUS": 8,
    "SEPTEMBER": 9,
    "OKTOBER": 10,
    "NOVEMBER": 11,
    "DESEMBER": 12,
}


def parse_lm_date(value) -> str | None:
    if value is None:
        return None
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    s = str(value).upper().strip()
    m = re.match(r"(\d{1,2})\s+(\w+)", s)
    if not m:
        return None
    month = MONTHS.get(m.group(2))
    if not month:
        return None
    year = 2026
    ym = re.search(r"(20\d{2})", s)
    if ym:
        year = int(ym.group(1))
    return f"{year:04d}-{month:02d}-{int(m.group(1)):02d}"


def section_end_row(ws, header_row: int) -> int:
    next_header: int | None = None
    for r in range(header_row + 1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v and str(v).strip().upper() == "HARGA ANTAM":
            next_header = r
            break
    if next_header is not None:
        end = next_header - 1
        while end > header_row:
            if any(ws.cell(end, c).value not in (None, "") for c in range(1, 9)):
                return end
            end -= 1
        return header_row + 11
    for r in range(ws.max_row, header_row, -1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, 9)):
            return r
    return header_row + 11


def parse_harga(val) -> int | None:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        v = float(val)
        return int(v) if v == int(v) else int(round(v))
    s = str(val).strip()
    if s in ("-", "—"):
        return None
    return None


def extract_ubs_rows(ws, header_row: int, end_row: int) -> list[tuple[float, int | None]]:
    ubs_header: int | None = None
    for r in range(header_row, end_row + 1):
        f = ws.cell(r, 6).value
        if f and str(f).strip().upper() == "HARGA UBS BATIK":
            ubs_header = r
            break
    if ubs_header is None:
        return []

    out: list[tuple[float, int | None]] = []
    for r in range(ubs_header + 1, end_row + 1):
        fv = ws.cell(r, 6).value
        gv = ws.cell(r, 7).value
        if fv is None:
            continue
        if isinstance(fv, str):
            fs = fv.strip()
            if fs.startswith("="):
                continue
            if not re.fullmatch(r"[\d.]+", fs):
                continue
            fv = float(fs)
        if isinstance(fv, (int, float)):
            out.append((float(fv), parse_harga(gv)))
    return out


def ensure_schema(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS harga_ubs_batik (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            tanggal TEXT NOT NULL,
            gram REAL NOT NULL,
            harga INTEGER,
            UNIQUE(tanggal, gram)
        )
        """
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_ubs_batik_tanggal ON harga_ubs_batik(tanggal)"
    )


def main() -> None:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb[SHEET]

    headers: list[int] = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v and str(v).strip().upper() == "HARGA ANTAM":
            headers.append(r)

    conn = sqlite3.connect(DB_PATH)
    ensure_schema(conn)

    imported_dates = 0
    total_rows = 0
    for h in headers:
        iso = parse_lm_date(ws.cell(h + 1, 1).value)
        if not iso:
            print(f"Lewati baris {h}: tanggal tidak dikenali ({ws.cell(h + 1, 1).value!r})")
            continue
        end = section_end_row(ws, h)
        rows = extract_ubs_rows(ws, h, end)
        if not rows:
            continue
        conn.execute("DELETE FROM harga_ubs_batik WHERE tanggal = ?", (iso,))
        conn.executemany(
            "INSERT INTO harga_ubs_batik (tanggal, gram, harga) VALUES (?, ?, ?)",
            [(iso, gram, harga) for gram, harga in rows],
        )
        print(f"{iso}: {len(rows)} baris UBS Batik")
        imported_dates += 1
        total_rows += len(rows)

    conn.commit()
    conn.close()
    wb.close()
    print(f"Selesai — {imported_dates} tanggal, {total_rows} baris → {DB_PATH}")


if __name__ == "__main__":
    main()
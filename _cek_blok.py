import openpyxl
p = r"D:\TOKO MAS LIA\DOKUMEN\TEMPLATE HARGA MAS2.xlsx"
wb = openpyxl.load_workbook(p)
ws = wb["ANTAM"]
starts = []
for r in range(1, ws.max_row + 1):
    v = ws.cell(r, 1).value
    if isinstance(v, str) and v.strip().upper() == "HARGA ANTAM":
        starts.append(r)
for s in starts[-6:]:
    tgl = ws.cell(s + 1, 1).value
    rows = []
    for r in range(s, s + 13):
        a = ws.cell(r, 1).value
        rows.append((a, ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 5).value))
    print("=== blok header baris", s, "| tgl:", tgl)
    for a, b, c, d, e in rows:
        print("   ", a, "| RETRO", b, "| RANDOM", c, "| 2025", d, "| 2026", e)

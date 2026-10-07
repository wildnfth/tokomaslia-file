#!/usr/bin/env python3
"""canva_mcp.py - klien MCP langsung ke mcp.canva.com (tanpa spawn agent LLM).

Pakai token OAuth yang sudah disimpan Hermes di %LOCALAPPDATA%/hermes/mcp-tokens/.
Refresh token otomatis kalau mau kedaluwarsa. Token TIDAK pernah dicetak.

Mode:
  tools                                  : daftar tool MCP
  read <design_id>                       : buka transaksi, dump teks mentah, cancel
  apply <design_id> --spec FILE          : replace teks sesuai spec JSON, commit, verifikasi
  sync  <design_id> --xlsx FILE          : sinkron blok terakhir Excel -> Canva (otomatis)

Struktur respons start-editing-transaction (yang relevan):
  transaction.transaction_id : id transaksi
  richtexts[]: {page_index, regions:[{type:"character", text}], element_id}
  pages[]: [{page_id, ...}]
"""
import argparse, json, os, re, sys, time, urllib.request, urllib.error, urllib.parse

HOME = os.environ.get("LOCALAPPDATA", os.path.expanduser("~") + "/AppData/Local")
TOKDIR = os.path.join(HOME, "hermes", "mcp-tokens")
URL = "https://mcp.canva.com/mcp"
PROTO = "2025-06-18"
INTENT = "Update harga logam mulia pada poster HARGA BELI sesuai Excel terbaru"
BULAN = {"JANUARI": "Januari", "FEBRUARI": "Februari", "MARET": "Maret", "APRIL": "April",
         "MEI": "Mei", "JUNI": "Juni", "JULI": "Juli", "AGUSTUS": "Agustus",
         "SEPTEMBER": "September", "OKTOBER": "Oktober", "NOVEMBER": "November",
         "DESEMBER": "Desember"}
RE_DATE = re.compile(r"^\d{1,2} (" + "|".join(BULAN.values()) + r") \d{4}$")

# ---------- token ----------

def _load(p):
    with open(p, encoding="utf-8") as f:
        return json.load(f)

def _save_atomic(path, data):
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    os.replace(tmp, path)

def get_token():
    tok = _load(os.path.join(TOKDIR, "canva.json"))
    if tok.get("expires_at", 0) > time.time() + 120:
        return tok["access_token"]
    meta = _load(os.path.join(TOKDIR, "canva.meta.json"))
    cli = _load(os.path.join(TOKDIR, "canva.client.json"))
    body = urllib.parse.urlencode({
        "grant_type": "refresh_token",
        "refresh_token": tok["refresh_token"],
        "client_id": cli["client_id"],
    }).encode()
    req = urllib.request.Request(meta["token_endpoint"], data=body,
                                 headers={"Content-Type": "application/x-www-form-urlencoded"})
    with urllib.request.urlopen(req, timeout=30) as r:
        j = json.load(r)
    tok["access_token"] = j["access_token"]
    if j.get("refresh_token"):
        tok["refresh_token"] = j["refresh_token"]
    tok["expires_in"] = j.get("expires_in", 14400)
    tok["expires_at"] = time.time() + tok["expires_in"] - 60
    _save_atomic(os.path.join(TOKDIR, "canva.json"), tok)
    return tok["access_token"]

# ---------- MCP over streamable HTTP ----------

class Mcp:
    def __init__(self):
        self.session = None
        self.next_id = 0

    def _post(self, payload):
        hdr = {
            "Content-Type": "application/json",
            "Accept": "application/json, text/event-stream",
            "Authorization": "Bearer " + get_token(),
            "MCP-Protocol-Version": PROTO,
        }
        if self.session:
            hdr["Mcp-Session-Id"] = self.session
        req = urllib.request.Request(URL, data=json.dumps(payload).encode(), headers=hdr)
        try:
            r = urllib.request.urlopen(req, timeout=60)
        except urllib.error.HTTPError as e:
            raise RuntimeError("HTTP %d: %s" % (e.code, e.read()[:400]))
        sid = r.headers.get("Mcp-Session-Id") or r.headers.get("mcp-session-id")
        if sid:
            self.session = sid
        ctype = r.headers.get("Content-Type", "")
        raw = r.read().decode("utf-8", "replace")
        if "text/event-stream" in ctype:
            msgs = []
            for line in raw.splitlines():
                if line.startswith("data:"):
                    data = line[5:].strip()
                    if data and data != "[DONE]":
                        try:
                            msgs.append(json.loads(data))
                        except json.JSONDecodeError:
                            pass
            return msgs
        if raw.strip():
            return [json.loads(raw)]
        return []

    def call(self, method, params=None, notify=False):
        self.next_id += 1
        payload = {"jsonrpc": "2.0", "method": method, "id": self.next_id}
        if params is not None:
            payload["params"] = params
        if notify:
            payload.pop("id")
        msgs = self._post(payload)
        if notify:
            return None
        for m in msgs:
            if m.get("id") == self.next_id:
                if "error" in m:
                    raise RuntimeError("MCP error: %s" % json.dumps(m["error"])[:400])
                return m.get("result")
        raise RuntimeError("tidak ada respons utk id %d (dapat %d pesan)" % (self.next_id, len(msgs)))

    def init(self):
        res = self.call("initialize", {
            "protocolVersion": PROTO,
            "capabilities": {},
            "clientInfo": {"name": "lia-canva-direct", "version": "1.0"},
        })
        self.call("notifications/initialized", {}, notify=True)
        return res

    def tool(self, name, args):
        res = self.call("tools/call", {"name": name, "arguments": args})
        if res is None:
            raise RuntimeError("hasil kosong")
        if res.get("isError"):
            raise RuntimeError("tool error: " + json.dumps(res)[:500])
        return res

def tool_result_text(res):
    out = []
    for c in res.get("content", []):
        if c.get("type") == "text":
            out.append(c.get("text", ""))
    if not out and "structuredContent" in res:
        out.append(json.dumps(res["structuredContent"], ensure_ascii=False))
    return "\n".join(out)

def tool_result_json(res):
    if res.get("structuredContent"):
        return res["structuredContent"]
    for c in res.get("content", []):
        if c.get("type") == "text":
            try:
                return json.loads(c["text"])
            except (json.JSONDecodeError, KeyError):
                pass
    return None

def find_tool(mcp, kw):
    res = mcp.call("tools/list")
    for t in res.get("tools", []):
        if kw in t["name"]:
            return t["name"]
    raise RuntimeError("tool dengan kata '%s' tidak ketemu: %s" %
                       (kw, [t["name"] for t in res.get("tools", [])]))

# ---------- struktur ----------

def parse_start(j):
    """(txn_id, texts{element_id: text}, pages_raw) dari respons start."""
    tr = j.get("transaction") or {}
    txn = tr.get("transaction_id") or tr.get("id")
    texts = {}
    for rt in j.get("richtexts", []):
        eid = rt.get("element_id")
        if not eid:
            continue
        parts = [r.get("text") for r in rt.get("regions", []) if isinstance(r.get("text"), str)]
        txt = "\n".join(p for p in parts if p.strip())
        if txt:
            texts[eid] = txt
    return txn, texts, j.get("pages", [])

def dump_texts(texts):
    for eid, txt in texts.items():
        print("  id=...%s" % eid[-14:])
        for ln in txt.splitlines():
            print("    %r" % ln)

def perform_commit_verify(mcp, design_id, ops, checks):
    """ops: [(type, element_id, text)]; checks: [(label, old_text, new_text)]."""
    start = find_tool(mcp, "start-editing")
    perform = find_tool(mcp, "perform-editing")
    commit = find_tool(mcp, "commit-editing")
    cancel = find_tool(mcp, "cancel-editing")

    res = mcp.tool(start, {"design_id": design_id})
    j = tool_result_json(res)
    if j is None:
        raise RuntimeError("start tidak mengembalikan JSON:\n" + tool_result_text(res)[:1500])
    txn, texts, pages_raw = parse_start(j)
    if not txn:
        raise RuntimeError("transaction_id tidak ketemu:\n" + json.dumps(j)[:800])
    print("txn:", txn, "| elemen ber-teks:", len(texts))

    final_ops = [{"type": typ, "element_id": eid, "text": txt} for typ, eid, txt in ops]
    args = {"transaction_id": txn, "operations": final_ops, "page_index": 1,
            "user_intent": INTENT}
    if pages_raw:
        args["pages"] = pages_raw
    print("== perform (%d operasi) ==" % len(final_ops))
    pres = mcp.tool(perform, args)
    pjson = tool_result_json(pres) or {}
    s = json.dumps(pjson).lower()
    if '"fail' in s or '"error"' in pjson:
        mcp.tool(cancel, {"transaction_id": txn})
        raise RuntimeError("ada operasi gagal, transaksi di-cancel: " + json.dumps(pjson)[:800])
    print("perform OK")

    print("== commit ==")
    cres = mcp.tool(commit, {"transaction_id": txn, "user_intent": INTENT})
    cjson = tool_result_json(cres) or {}
    print("committed:", json.dumps(cjson)[:160] if cjson else tool_result_text(cres)[:160])

    print("== verifikasi ulang (transaksi baru, lalu cancel) ==")
    res2 = mcp.tool(start, {"design_id": design_id})
    j2 = tool_result_json(res2)
    txn2, texts2, _ = parse_start(j2 or {})
    ok = True
    for label, old_t, new_t in checks:
        n_old = sum(1 for t in texts2.values() if t == old_t)
        n_new = sum(1 for t in texts2.values() if t == new_t)
        # kotak yang memang tidak berubah (old == new): cukup pastikan teksnya ada
        good = (n_new >= 1) if old_t == new_t else (n_old == 0 and n_new >= 1)
        print("  %-12s lama tersisa=%d baru=%d %s" % (label, n_old, n_new, "OK" if good else "GAGAL"))
        ok = ok and good
    if txn2:
        mcp.tool(cancel, {"transaction_id": txn2})
    if not ok:
        raise RuntimeError("verifikasi gagal: masih ada teks lama / teks baru tidak ada")
    print("SEMUA OK")

# ---------- mode read / apply ----------

def cmd_read(mcp, design_id):
    start = find_tool(mcp, "start-editing")
    cancel = find_tool(mcp, "cancel-editing")
    res = mcp.tool(start, {"design_id": design_id})
    j = tool_result_json(res)
    if j is None:
        print(tool_result_text(res))
        return
    txn, texts, _ = parse_start(j)
    print("transaction_id:", txn)
    print("=== teks mentah semua elemen ===")
    dump_texts(texts)
    if txn:
        mcp.tool(cancel, {"transaction_id": txn})
        print("cancelled:", txn)


def cmd_apply(mcp, design_id, spec_path):
    spec = json.load(open(spec_path, encoding="utf-8"))
    start = find_tool(mcp, "start-editing")
    res = mcp.tool(start, {"design_id": design_id})
    j = tool_result_json(res)
    if j is None:
        raise RuntimeError("start tidak mengembalikan JSON:\n" + tool_result_text(res)[:1500])
    txn, texts, _ = parse_start(j)
    if not txn:
        raise RuntimeError("transaction_id tidak ketemu:\n" + json.dumps(j)[:800])
    print("txn:", txn, "| elemen ber-teks:", len(texts))

    used = set()
    ops, checks = [], []

    def match_all(old):
        return [eid for eid, txt in texts.items() if txt == old and eid not in used]

    for box in spec["replacements"]:
        hits = match_all(box["old_text"])
        if not hits:
            raise RuntimeError("[%s] old_text tidak ketemu persis di desain" % box["label"])
        for eid in hits:
            used.add(eid)
            ops.append(("replace_text", eid, box["new_text"]))
            print("  match %-12s id=...%s" % (box["label"], eid[-12:]))
        checks.append((box["label"], box["old_text"], box["new_text"]))
    if spec.get("replacements_date"):
        d = spec["replacements_date"]
        hits = match_all(d["old_text"])
        if hits:
            for eid in hits:
                used.add(eid)
                ops.append(("replace_text", eid, d["new_text"]))
                print("  match tanggal      id=...%s" % eid[-12:])
            checks.append(("tanggal", d["old_text"], d["new_text"]))
        else:
            print("  (tanggal tidak ketemu untuk diganti; lewati)")
    mcp.tool(find_tool(mcp, "cancel-editing"), {"transaction_id": txn})
    perform_commit_verify(mcp, design_id, ops, checks)

# ---------- mode sync ----------

def fmt_rp(v):
    return "{:,}".format(int(v)).replace(",", ".")

def read_excel_block(xlsx):
    """Baca blok terakhir sheet ANTAM -> (date_text, maps{brand:{gram: rupiah_str}}).
    Kolom: A gram-RETRO, B RETRO, C RANDOM(=ANTAM di Canva); G gram-G24, H G24;
    K gram-UBS, L UBS. Baris = header+2 .. header+11."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, data_only=True)
    if "ANTAM" not in wb.sheetnames:
        raise RuntimeError("sheet ANTAM tidak ada; sheets: %s" % wb.sheetnames)
    ws = wb["ANTAM"]
    hs = [r for r in range(1, ws.max_row + 1)
          if isinstance(ws.cell(r, 1).value, str) and ws.cell(r, 1).value.strip().upper().startswith("HARGA ANTAM")]
    if not hs:
        raise RuntimeError("blok HARGA ANTAM tidak ketemu di %s" % xlsx)
    h = hs[-1]
    raw_date = str(ws.cell(h + 1, 1).value or "").strip()
    parts = raw_date.split()
    date_canva = " ".join([parts[0], BULAN.get(parts[1].upper(), parts[1]), parts[2]]) if len(parts) == 3 else raw_date

    def val(v):
        if v is None:
            return None
        s = str(v).strip()
        if s in ("-", "SOLD"):
            return "SOLD"
        try:
            # Excel menyimpan harga dalam RIBUAN (1260 = Rp 1.260.000)
            return fmt_rp(float(s) * 1000)
        except ValueError:
            return None

    retro, random_, g24, ubs = {}, {}, {}, {}
    for r in range(h + 2, h + 12):
        gA = ws.cell(r, 1).value
        if isinstance(gA, (int, float)):
            b, c = val(ws.cell(r, 2).value), val(ws.cell(r, 3).value)
            if b: retro[round(float(gA), 1)] = b
            if c: random_[round(float(gA), 1)] = c
        gG = ws.cell(r, 7).value
        if isinstance(gG, (int, float)):
            vH = val(ws.cell(r, 8).value)
            if vH: g24[round(float(gG), 1)] = vH
        gK = ws.cell(r, 11).value
        if isinstance(gK, (int, float)):
            vL = val(ws.cell(r, 12).value)
            if vL: ubs[round(float(gK), 1)] = vL
    return date_canva, {"RETRO": retro, "ANTAM": random_, "GALERI24": g24, "UBS": ubs}

def identify_box(text):
    """Tebak kotak harga mana dari isi teksnya. Return nama brand atau None."""
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(lines) < 5 or not all(re.match(r"^\d+(\.\d+)? gr : ", ln) for ln in lines):
        return None
    if any(ln.startswith("4 gr : ") for ln in lines):
        return "UBS"
    if len(lines) == 7 and lines[0].startswith("1 gr : "):
        return "GALERI24"
    if len(lines) == 9 and lines[0].startswith("0.5 gr : "):
        return "RETRO" if lines[-1].startswith("100 gr : SOLD") else "ANTAM"
    return None

def box_from_map(mapping, labels):
    """Susun teks kotak dari map {gram: rupiah_str}, ikut urutan label yang ada."""
    out = []
    for lab in labels:
        g = float(lab.split(" ")[0])
        v = mapping.get(round(g, 1))
        if v is None:
            return None
        out.append("%s : %s" % (lab.split(" : ")[0], v))
    return "\n".join(out)

def cmd_sync(mcp, design_id, xlsx):
    start = find_tool(mcp, "start-editing")
    cancel = find_tool(mcp, "cancel-editing")
    date_canva, maps = read_excel_block(xlsx)
    print("Excel blok terakhir:", date_canva)

    res = mcp.tool(start, {"design_id": design_id})
    j = tool_result_json(res)
    txn, texts, _ = parse_start(j)
    if not txn:
        raise RuntimeError("transaction_id tidak ketemu")
    print("txn:", txn, "| elemen ber-teks:", len(texts))

    used = set()
    ops, checks = [], []
    found = {}
    for eid, txt in texts.items():
        brand = identify_box(txt)
        if brand:
            found.setdefault(brand, []).append((eid, txt))

    for brand in ("ANTAM", "RETRO", "GALERI24", "UBS"):
        elems = found.get(brand, [])
        if not elems:
            raise RuntimeError("kotak %s tidak teridentifikasi di desain" % brand)
        sample = elems[0][1]
        labels = [ln.split(" : ")[0] for ln in sample.splitlines()]
        want = box_from_map(maps[brand], labels)
        if want is None:
            raise RuntimeError("nilai Excel %s kurang utk label %s" % (brand, labels))
        for eid, txt in elems:
            used.add(eid)
            if txt != want:
                ops.append(("replace_text", eid, want))
                print("  %-9s id=...%s DIGANTI" % (brand, eid[-12:]))
            else:
                print("  %-9s id=...%s sama" % (brand, eid[-12:]))
        checks.append((brand, sample, want))

    for eid, txt in texts.items():
        if eid in used or not RE_DATE.match(txt.strip()):
            continue
        used.add(eid)
        if txt.strip() != date_canva:
            ops.append(("replace_text", eid, date_canva))
            print("  tanggal  id=...%s %r -> %r DIGANTI" % (eid[-12:], txt, date_canva))
        else:
            print("  tanggal  id=...%s sama" % eid[-12:])

    if not ops:
        mcp.tool(cancel, {"transaction_id": txn})
        print("Canva sudah sinkron dengan Excel. Tidak ada perubahan (transaksi di-cancel).")
        return
    mcp.tool(cancel, {"transaction_id": txn})
    perform_commit_verify(mcp, design_id, ops, checks)

# ---------- main ----------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("mode", choices=["tools", "read", "apply", "sync"])
    ap.add_argument("design_id")
    ap.add_argument("--spec", help="file JSON spec utk mode apply")
    ap.add_argument("--xlsx", help="file Excel utk mode sync")
    a = ap.parse_args()
    mcp = Mcp()
    mcp.init()
    if a.mode == "tools":
        res = mcp.call("tools/list")
        for t in res.get("tools", []):
            print(t["name"])
    elif a.mode == "read":
        cmd_read(mcp, a.design_id)
    elif a.mode == "apply":
        if not a.spec:
            ap.error("--spec wajib utk apply")
        cmd_apply(mcp, a.design_id, a.spec)
    elif a.mode == "sync":
        if not a.xlsx:
            ap.error("--xlsx wajib utk sync")
        cmd_sync(mcp, a.design_id, a.xlsx)


if __name__ == "__main__":
    main()

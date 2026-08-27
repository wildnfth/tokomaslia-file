# -*- coding: utf-8 -*-
from openpyxl import load_workbook

path = r"D:\TOKO MAS LIA\DOKUMEN\TEMPLATE HARGA MAS2.xlsx"
wb = load_workbook(path, data_only=False)
ws = wb["ANTAM"]

src_h, dst_h = 1487, 1500
src_end, dst_end = 1498, 1511
n_rows = src_end - src_h + 1

print("=== dates ===")
for c, name in ((1, "A"), (7, "G"), (11, "K")):
    print(name, dst_h, ws.cell(dst_h, c).value)
    print(name, dst_h + 1, ws.cell(dst_h + 1, c).value)

print("=== prices dst (not multiple of 5?) ===")
bad = []
price_cols = {2, 3, 4, 5, 8, 12}
for r in range(dst_h, dst_end + 1):
    row = []
    for c in range(1, 14):
        v = ws.cell(r, c).value
        if c in price_cols and isinstance(v, (int, float)) and not isinstance(v, bool):
            if v % 5 != 0:
                bad.append((r, c, v))
        if v not in (None, ""):
            row.append("%s=%s" % (chr(64 + c) if c <= 26 else c, v))
    if row:
        print(r, " | ".join(row))
print("bad_not_mult5", bad)

print("=== merged offset ===")
src_merged = sorted(str(m) for m in ws.merged_cells.ranges if not (m.max_row < src_h or m.min_row > src_end))
dst_merged = sorted(str(m) for m in ws.merged_cells.ranges if not (m.max_row < dst_h or m.min_row > dst_end))
offset = dst_h - src_h
print("src", src_merged)
print("dst", dst_merged)
print("count src/dst", len(src_merged), len(dst_merged))

# check offset match roughly
from openpyxl.utils import range_boundaries
mismatch = 0
for s in src_merged:
    min_col, min_row, max_col, max_row = range_boundaries(s)
    expected = None
    # rebuild expected by shifting rows
    from openpyxl.utils import get_column_letter
    exp = "%s%s:%s%s" % (
        get_column_letter(min_col), min_row + offset,
        get_column_letter(max_col), max_row + offset,
    )
    if exp not in dst_merged:
        mismatch += 1
        print("MISSING merge", s, "->", exp)
print("merge_mismatch", mismatch)

print("=== border / number_format sample ===")
for r_off in (0, 1, 3, n_rows - 1):
    for c in (1, 2, 8, 12):
        src = ws.cell(src_h + r_off, c)
        dst = ws.cell(dst_h + r_off, c)
        sb = src.border.left.style if src.border and src.border.left else None
        db = dst.border.left.style if dst.border and dst.border.left else None
        print("r+%d c%d src_fmt=%s dst_fmt=%s src_b=%s dst_b=%s" % (
            r_off, c, src.number_format, dst.number_format, sb, db))

print("=== 0.5g ANTAM snap check (should be 1350 from 1347.5) ===")
print("dst 0.5 RETRO B", ws.cell(1503, 2).value)
print("dst 0.5 RANDOM C", ws.cell(1503, 3).value)
print("dst UBS 0.5 L", ws.cell(1502, 12).value)

wb.close()

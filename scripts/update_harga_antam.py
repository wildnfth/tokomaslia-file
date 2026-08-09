#!/usr/bin/env python3
"""
Update HARGA ANTAM (sheet ANTAM) — TAMBAH blok harian baru.
Pola: salin blok harga terakhir, tambah 1 baris kosong, update harga naik/turun
sesuai --step, ganti tanggal, pertahankan format sel MERAH = KOSONG.

Jalankan:
    python update_harga_antam.py "TEMPLATE HARGA MAS2.xlsx" --date "8 AGUSTUS 2026" --step -40
    python update_harga_antam.py "TEMPLATE HARGA MAS2.xlsx" --date "8 AGUSTUS 2026" --step -40 --dry-run
"""

import argparse
import os
import shutil
import sys
from copy import copy
from datetime import datetime

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl tidak terinstal. Jalankan: pip install openpyxl")
    sys.exit(1)

def parse_args():
    parser = argparse.ArgumentParser(description="Update harga ANTAM (tambah blok baru)")
    parser.add_argument("file", help="Path ke file Excel TEMPLATE HARGA MAS2.xlsx")
    parser.add_argument("--date", required=True, help="Tanggal baru, contoh: '8 AGUSTUS 2026'")
    parser.add_argument("--step", type=int, default=50, help="Kenaikan/turun per gram (default 50)")
    parser.add_argument("--dry-run", action="store_true", help="Tampilkan preview tanpa menulis")
    parser.add_argument("--no-backup", action="store_true", help="Nonaktifkan backup otomatis")
    return parser.parse_args()


def backup_file(filepath):
    """Buat backup file dengan timestamp."""
    if not os.path.exists(filepath):
        return None
    base, ext = os.path.splitext(filepath)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = f"{base}_BACKUP_{timestamp}{ext}"
    shutil.copy2(filepath, backup_path)
    print(f"[BACKUP] {filepath} -> {backup_path}")
    return backup_path

def find_last_block(ws):
    """Cari blok terakhir di sheet ANTAM (era baru)."""
    max_row = ws.max_row
    for r in range(max_row, 1, -1):
        cell_a = ws.cell(row=r, column=1).value
        if cell_a and isinstance(cell_a, str) and "HARGA ANTAM" in cell_a and "-" in cell_a:
            metadata_row = r + 1
            data_start = r + 2
            data_end = data_start
            for dr in range(data_start, max_row + 1):
                val_a = ws.cell(row=dr, column=1).value
                val_g = ws.cell(row=dr, column=7).value
                if val_a is None and val_g is None:
                    break
                if val_a == "MERAH = KOSONG" or val_g == "MERAH = KOSONG":
                    next_a = ws.cell(row=dr + 1, column=1).value if dr + 1 <= max_row else None
                    next_g = ws.cell(row=dr + 1, column=7).value if dr + 1 <= max_row else None
                    if next_a is None and next_g is None:
                        data_end = dr
                        break
                data_end = dr
            return r, data_start, data_end, metadata_row
    raise ValueError("Tidak dapat menemukan blok HARGA ANTAM terakhir di sheet.")

def copy_style(src, dst):
    """Salin style sel (font, fill, border, format angka, alignment, protection)."""
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.number_format = src.number_format
    dst.alignment = copy(src.alignment)
    dst.protection = copy(src.protection)


def copy_block(ws, header_row, data_start, data_end, metadata_row, new_date, step, dry_run=False):
    """Salin blok terakhir (header + metadata + data) ke bawah BESERTA FORMAT
    (style sel, merged cells, tinggi baris), update harga & tanggal.
    Hasilnya plek ketiplek dengan blok sumber."""
    new_header_row = data_end + 2
    new_metadata_row = new_header_row + 1
    new_data_start = new_metadata_row + 1
    delta = new_header_row - header_row

    new_header_a = f"HARGA ANTAM - {new_date}"
    new_header_g = "HARGA Galeri24"
    new_header_k = "HARGA UBS BATIK"

    if dry_run:
        print(f"[DRY-RUN] Akan menambahkan blok baru di baris {new_header_row} (COPAS nilai+style+merge+tinggi)")

    # 1) Salin tinggi baris header + metadata + data
    for r in range(header_row, data_end + 1):
        rd = ws.row_dimensions.get(r)
        if rd is not None and rd.height:
            if dry_run:
                print(f"[DRY-RUN] Baris {r + delta}: tinggi baris = {rd.height}")
            else:
                ws.row_dimensions[r + delta].height = rd.height

    # 2) Kumpulkan merged cells pada area blok sumber (untuk diterapkan di blok baru)
    merges_src = []
    for mr in ws.merged_cells.ranges:
        if header_row <= mr.min_row <= data_end:
            merges_src.append((mr.min_row + delta, mr.min_col, mr.max_row + delta, mr.max_col))

    # 3) Salin SEMUA sel (nilai + style) — col 1..16
    total_rows = 0
    for r in range(header_row, data_end + 1):
        new_r = r + delta
        for col in range(1, 17):
            src = ws.cell(r, col)
            if dry_run:
                if src.value is not None:
                    print(f"[DRY-RUN] {new_r} {get_column_letter(col)}: salin '{src.value}'")
            else:
                dst = ws.cell(new_r, col)
                copy_style(src, dst)
                dst.value = src.value
        total_rows += 1

    # 4) Update tanggal di header & metadata
    if dry_run:
        print(f"[DRY-RUN] Header A{new_header_row} -> '{new_header_a}'")
        print(f"[DRY-RUN] Metadata G{new_metadata_row} & K{new_metadata_row}: tanggal -> '{new_date}'")
    else:
        ws.cell(new_header_row, 1, new_header_a)
        ws.cell(new_header_row, 7, new_header_g)
        ws.cell(new_header_row, 11, new_header_k)
        ws.cell(new_metadata_row, 7, new_date)
        ws.cell(new_metadata_row, 11, new_date)

    # 5) Update harga per gram sesuai step (lewati baris MERAH = KOSONG)
    price_cols = {2: "B", 3: "C", 4: "D", 5: "E", 8: "H", 9: "I", 12: "L"}
    for i, old_row in enumerate(range(data_start, data_end + 1)):
        new_row = new_data_start + i
        val_a = ws.cell(old_row, 1).value
        val_g = ws.cell(old_row, 7).value
        if val_a == "MERAH = KOSONG" or val_g == "MERAH = KOSONG":
            merah_col = 7 if val_g == "MERAH = KOSONG" else 1
            if dry_run:
                print(f"[DRY-RUN] Baris {new_row}: pertahankan MERAH = KOSONG (kol {merah_col})")
            continue
        for col, desc in price_cols.items():
            val = ws.cell(old_row, col).value
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                new_val = val + step
                if dry_run:
                    print(f"[DRY-RUN] {new_row} {desc}: {val:,.0f} -> {new_val:,.0f}")
                else:
                    ws.cell(new_row, col, new_val)

    # 6) Terapkan merged cells pada blok baru (hindari duplikat)
    if not dry_run:
        existing = {(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges}
        for (mr, mc, mrr, mrc) in merges_src:
            if (mr, mc, mrr, mrc) not in existing:
                ws.merge_cells(start_row=mr, start_column=mc, end_row=mrr, end_column=mrc)

    print(f"[INFO] Total baris blok disalin: {total_rows} (dengan style, merge, tinggi baris)")
    return new_header_row


def main():
    args = parse_args()
    filepath = os.path.abspath(args.file)

    if not os.path.exists(filepath):
        print(f"ERROR: File tidak ditemukan: {filepath}")
        sys.exit(1)

    print(f"[INFO] File: {filepath}")
    print(f"[INFO] Tanggal: {args.date}, Step: {args.step:+d}, Dry-run: {args.dry_run}")

    wb = load_workbook(filepath)
    if "ANTAM" not in wb.sheetnames:
        print("ERROR: Sheet 'ANTAM' tidak ditemukan.")
        sys.exit(1)
    ws = wb["ANTAM"]

    header_row, data_start, data_end, metadata_row = find_last_block(ws)
    print(f"[INFO] Blok terakhir: header={header_row}, metadata={metadata_row}, data={data_start}-{data_end}")

    if args.dry_run:
        print("\n" + "=" * 60)
        print("DRY-RUN MODE — Tidak ada yang ditulis ke file")
        print("=" * 60)
        copy_block(ws, header_row, data_start, data_end, metadata_row, args.date, args.step, dry_run=True)
        print("\n[INFO] Jika sesuai, jalankan lagi tanpa --dry-run")
        return

    if not args.no_backup:
        backup_file(filepath)

    new_header_row = copy_block(ws, header_row, data_start, data_end, metadata_row, args.date, args.step, dry_run=False)
    wb.save(filepath)
    print(f"\n[SUCCESS] {filepath} diupdate, blok baru di baris {new_header_row}")


if __name__ == "__main__":
    main()


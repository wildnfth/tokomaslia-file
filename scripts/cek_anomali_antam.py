# -*- coding: utf-8 -*-
"""Cek urutan harga ANTAM: 2026 > 2025 > RANDOM. RETRO diabaikan. Tidak menulis file."""
import argparse
import sys
from datetime import datetime, timedelta

from openpyxl import load_workbook

HEADER_MARK = "HARGA ANTAM"
MONTHS = {
    "JANUARI": 1, "FEBRUARI": 2, "MARET": 3, "APRIL": 4, "MEI": 5, "JUNI": 6,
    "JULI": 7, "AGUSTUS": 8, "SEPTEMBER": 9, "OKTOBER": 10, "NOVEMBER": 11, "DESEMBER": 12,
}


def to_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def lab(v):
    if v is None or v == "":
        return ""
    return str(v).strip().upper()


def parse_date(s):
    if not s:
        return None
    parts = s.strip().upper().replace(",", " ").split()
    # "28 AGUSTUS 2026" or "14 APRIL SIANG"
    if len(parts) < 2:
        return None
    try:
        day = int(parts[0])
    except ValueError:
        return None
    mon = MONTHS.get(parts[1])
    if not mon:
        return None
    year = None
    for p in parts[2:]:
        if p.isdigit() and len(p) == 4:
            year = int(p)
            break
    if year is None:
        year = datetime.now().year
    try:
        return datetime(year, mon, day)
    except ValueError:
        return None


def find_headers(ws):
    rows = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip().startswith(HEADER_MARK):
            rows.append(r)
    return rows


def block_meta(ws, headers, i):
    h = headers[i]
    end = (headers[i + 1] - 1) if i + 1 < len(headers) else ws.max_row
    date_val = None
    for r in range(h, min(h + 4, end + 1)):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip() and not v.strip().startswith(HEADER_MARK):
            date_val = v.strip()
            break
    label_row = None
    cols = {}
    for r in range(h, min(h + 6, end + 1)):
        found = {}
        for c in range(2, 6):
            t = lab(ws.cell(r, c).value)
            if t:
                found[t] = c
        if "RANDOM" in found and "2025" in found and "2026" in found:
            label_row = r
            cols = found
            break
    return {
        "header": h,
        "end": end,
        "date": date_val or ("header r%d" % h),
        "parsed": parse_date(date_val or ""),
        "label_row": label_row,
        "cols": cols,
    }


def scan_block(ws, meta):
    if not meta["label_row"]:
        return []
    c_rand = meta["cols"]["RANDOM"]
    c_25 = meta["cols"]["2025"]
    c_26 = meta["cols"]["2026"]
    issues = []
    for r in range(meta["label_row"] + 1, meta["end"] + 1):
        gram = to_num(ws.cell(r, 1).value)
        if gram is None or not (0.4 <= gram <= 100):
            continue
        rand = to_num(ws.cell(r, c_rand).value)
        y25 = to_num(ws.cell(r, c_25).value)
        y26 = to_num(ws.cell(r, c_26).value)
        if rand is None or y25 is None or y26 is None:
            continue
        parts = []
        if not (y26 > y25):
            rel = "=" if y26 == y25 else "<"
            parts.append("2026%s2025 selisih %+g" % (rel, y26 - y25))
        if not (y25 > rand):
            rel = "=" if y25 == rand else "<"
            parts.append("2025%sRANDOM selisih %+g" % (rel, y25 - rand))
        if not (y26 > rand):
            rel = "=" if y26 == rand else "<"
            parts.append("2026%sRANDOM selisih %+g" % (rel, y26 - rand))
        if parts:
            issues.append({
                "row": r,
                "gram": gram,
                "random": rand,
                "y25": y25,
                "y26": y26,
                "parts": parts,
            })
    return issues


def select_blocks(metas, days=7, last=None, all_blocks=False):
    complete = [m for m in metas if m["label_row"]]
    if all_blocks:
        return complete
    if last is not None:
        n = max(1, int(last))
        return complete[-n:]
    dated = [m for m in complete if m["parsed"]]
    if not dated:
        return complete[-8:]
    latest = max(m["parsed"] for m in dated)
    cutoff = latest - timedelta(days=days)
    picked = [m for m in complete if m["parsed"] and m["parsed"] >= cutoff]
    return picked or complete[-8:]


def run_check(path, days=7, last=None, all_blocks=False):
    wb = load_workbook(path, data_only=True)
    if "ANTAM" not in wb.sheetnames:
        print("[err] sheet ANTAM tidak ada")
        return 1
    ws = wb["ANTAM"]
    headers = find_headers(ws)
    metas = [block_meta(ws, headers, i) for i in range(len(headers))]
    chosen = select_blocks(metas, days=days, last=last, all_blocks=all_blocks)
    if not chosen:
        print("[ok] tidak ada blok ANTAM dengan kolom RANDOM/2025/2026")
        return 0
    print("aturan: 2026 > 2025 > RANDOM  (RETRO diabaikan)")
    print("blok dicek:", len(chosen), "| dari", chosen[0]["date"], "s/d", chosen[-1]["date"])
    total = 0
    for m in chosen:
        issues = scan_block(ws, m)
        if not issues:
            print("[ok]", m["date"])
            continue
        total += len(issues)
        print("[ANOMALI]", m["date"], "(%d gramasi)" % len(issues))
        for it in issues:
            print(
                "    %g g  RANDOM=%s  2025=%s  2026=%s  -> %s"
                % (it["gram"], _fmt(it["random"]), _fmt(it["y25"]), _fmt(it["y26"]), "; ".join(it["parts"]))
            )
    print("-" * 40)
    if total:
        print("TOTAL ANOMALI:", total, "| jangan Discord, laporkan dulu, jangan fix tanpa izin")
        return 1
    print("TOTAL ANOMALI: 0")
    return 0


def _fmt(v):
    if v is None:
        return "-"
    if float(v).is_integer():
        return str(int(v))
    return str(v)


def main():
    ap = argparse.ArgumentParser(description="Cek urutan ANTAM 2026>2025>RANDOM.")
    ap.add_argument("file", nargs="?", default="TEMPLATE HARGA MAS2.xlsx")
    ap.add_argument("--days", type=int, default=7, help="Rentang hari dari blok terbaru (default 7).")
    ap.add_argument("--last", type=int, default=None, help="Hanya N blok terakhir.")
    ap.add_argument("--all", action="store_true", help="Semua blok yang punya 3 kolom.")
    args = ap.parse_args()
    code = run_check(args.file, days=args.days, last=args.last, all_blocks=args.all)
    sys.exit(code)


if __name__ == "__main__":
    main()

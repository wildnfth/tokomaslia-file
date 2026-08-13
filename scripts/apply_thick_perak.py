# -*- coding: utf-8 -*-
"""Terapkan border THICK pada blok terbaru TEMPLATE HARGA PERAK.xlsx."""
import sys
from openpyxl import load_workbook
from openpyxl.styles import Border, Side

FILE = sys.argv[1] if len(sys.argv) > 1 else r"D:/TOKO MAS LIA/DOKUMEN/TEMPLATE HARGA PERAK.xlsx"
THICK = Side(style="thick", color="000000")


from update_harga_perak import apply_thick_borders


def main():
    wb = load_workbook(FILE)
    ws = wb.active
    # cari blok terakhir (baris TANGGAL)
    starts = [r for r in range(1, ws.max_row + 1)
              if isinstance(ws.cell(r, 1).value, str)
              and ws.cell(r, 1).value.strip().upper().startswith("TANGGAL ")]
    if not starts:
        raise SystemExit('[err] tidak ada baris TANGGAL di kolom A.')
    start = starts[-1]

    apply_thick_borders(ws, start)

    wb.save(FILE)
    print("[ok] thick border diterapkan pada blok %d (%s)"
          % (start, ws.cell(start, 1).value))


if __name__ == "__main__":
    main()

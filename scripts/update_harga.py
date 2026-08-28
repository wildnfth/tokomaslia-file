# -*- coding: utf-8 -*-
"""
update_harga.py - Satu script untuk update harga (ANTAM + Galeri24 + UBS + EMAS).

Format blok ANTAM BARU (2 baris header) sejak blok 24 AGUSTUS 2026:
  row A  : 'HARGA ANTAM' | G : 'HARGA Galeri24' | K : 'HARGA UBS BATIK'
  row B  : 'TANGGAL'     | G : 'TANGGAL'        | K : 'TANGGAL'
  row C  : label merk ANTAM (RETRO/RANDOM/2025/2026) + Galeri24 gr1 + UBS gr0.5
  row D+ : harga ANTAM mulai (gram di kolom A)  + Galeri24/UBS harga

MODE:
  add    (default) : TAMBAH blok harian baru (menyalin utuh blok terakhir) utk ANTAM.
  update           : UBAH harga blok terakhir DI TEMPAT (tanpa blok baru).
  json             : update sheet EMAS (perhiasan) dari file JSON.

CONTOH:
  update_harga.py <file> --date "25 AGUSTUS 2026" --step 10 --targets allam
  update_harga.py <file> --date "25 AGUSTUS 2026" --step 10            # ANTAM aja
  update_harga.py <file> --mode update --targets ubs --step 30
  update_harga.py <file> --mode update --targets galeri24 --step 25
  update_harga.py <file> --mode update --grams "1,5" --step 40
  update_harga.py <file> --json perhiasan.json
  # selalu mulai dgn --dry-run
"""
import argparse
import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime

from openpyxl import load_workbook

HEADER_MARK = "HARGA ANTAM"
GIT = r"C:\Program Files\Git\cmd\git.exe"
REPO = os.path.abspath(os.getcwd())

GRAM_COL = {2: 1, 3: 1, 4: 1, 5: 1, 8: 7, 12: 11}
TARGET_MAP = {
    "retro": [2], "random": [3], "2025": [4], "2026": [5],
    "galeri24": [8], "ubs": [12],
}
ANTAM_NAMES = ["retro", "random", "2025", "2026"]


def find_header_rows(ws):
    rows = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip().startswith(HEADER_MARK):
            rows.append(r)
    return rows


def copy_style(src, dst):
    if src.has_style:
        dst._style = copy.copy(src._style)


def resolve_columns(spec):
    cols, labels = [], set()
    if not spec:
        spec = "antam"
    for t in spec.split(","):
        t = t.strip().lower()
        if t in ("allam", "all", "semua"):
            for name in ANTAM_NAMES:
                labels.add(name); cols.extend(TARGET_MAP[name])
            labels.add("galeri24"); cols.extend(TARGET_MAP["galeri24"])
            labels.add("ubs"); cols.extend(TARGET_MAP["ubs"])
        elif t == "antam":
            for name in ANTAM_NAMES:
                labels.add(name); cols.extend(TARGET_MAP[name])
        elif t in TARGET_MAP:
            labels.add(t); cols.extend(TARGET_MAP[t])
        else:
            print("[warn] target tidak dikenal, dilewati:", t)
    cols = list(dict.fromkeys(cols))
    return cols, labels


def resolve_grams(spec):
    if not spec:
        return None
    out = set()
    for g in spec.split(","):
        try:
            out.add(float(g.strip()))
        except ValueError:
            print("[warn] gramasi tidak valid:", g)
    return out or None


def snap5(v):
    return int(round(v / 5.0) * 5)


def last_data_row(ws, start):
    """Baris terakhir yang masih berisi nilai di kolom A-M, dari header blok."""
    last = start
    for r in range(start, ws.max_row + 1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, 14)):
            last = r
    return last


def unmerge_overlap(ws, start, end):
    doomed = []
    for m in list(ws.merged_cells.ranges):
        if not (m.max_row < start or m.min_row > end):
            doomed.append(str(m))
    for addr in doomed:
        ws.unmerge_cells(addr)


def apply_inplace(ws, wsv, rows, step, selected_cols, gram_filter):
    changed = []
    for r in rows:
        for c in selected_cols:
            gsrc = ws.cell(r, GRAM_COL[c]).value
            try:
                g = float(gsrc)
            except (TypeError, ValueError):
                continue
            if gram_filter is not None and g not in gram_filter:
                continue
            v = ws.cell(r, c).value
            base = v
            if isinstance(v, str) and v.startswith("="):
                base = wsv.cell(r, c).value
            try:
                base = float(base)
            except (TypeError, ValueError):
                continue
            nv = base + step * g
            changed.append((r, c, g, base, nv))
            ws.cell(r, c).value = nv
    return changed


def apply_round5(ws, rows, selected_cols):
    n = 0
    for r in rows:
        for c in selected_cols:
            gsrc = ws.cell(r, GRAM_COL[c]).value
            try:
                float(gsrc)
            except (TypeError, ValueError):
                continue
            cell = ws.cell(r, c)
            v = cell.value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                nv = snap5(v)
                if nv != v:
                    cell.value = nv
                    n += 1
    return n


def do_add(ws, wb, wsv, step, selected_cols, gram_filter, new_date, dry_run, outfile):
    headers = find_header_rows(ws)
    if not headers:
        sys.exit("Tidak ada blok 'HARGA ANTAM'.")
    src_h = headers[-1]
    src_end = last_data_row(ws, src_h)
    block_rows = list(range(src_h, src_end + 1))
    dst_h = src_end + 2
    spacing = dst_h - src_h
    n_rows = len(block_rows)
    date_row = src_h + 1
    print("[info] step %d/gram | target %s" % (step, sorted(selected_cols)))
    print("[info] blok sumber: %s-%s | tanggal lama: %s" % (src_h, src_end, ws.cell(date_row, 1).value))
    if dry_run:
        print("[dry-run/add] blok baru: %s-%s -> header '%s' / tanggal '%s'" % (
            dst_h, dst_h + n_rows - 1, HEADER_MARK, new_date))
        for di, sr in enumerate(block_rows):
            dr = dst_h + di
            for c in selected_cols:
                g = ws.cell(sr, GRAM_COL[c]).value
                try:
                    g = float(g)
                except (TypeError, ValueError):
                    continue
                if gram_filter is not None and g not in gram_filter:
                    continue
                v = ws.cell(sr, c).value
                try:
                    base = float(v)
                except (TypeError, ValueError):
                    continue
                print("  dr %d | gram %s | col %s | %s -> %s" % (dr, g, c, base, base + step * g))
        print("[dry-run] selesai (belum disimpan).")
        return

    dst_end = dst_h + n_rows - 1
    src_merges = [m for m in list(ws.merged_cells.ranges)
                  if not (m.max_row < src_h or m.min_row > src_end)]
    unmerge_overlap(ws, src_end + 1, dst_end)
    dst_rows = list(range(dst_h, dst_end + 1))
    for di, sr in enumerate(block_rows):
        dr = dst_rows[di]
        for c in range(1, ws.max_column + 1):
            s = ws.cell(sr, c)
            d = ws.cell(dr, c)
            d.value = s.value
            copy_style(s, d)
    for di, sr in enumerate(block_rows):
        dr = dst_rows[di]
        if sr in ws.row_dimensions and ws.row_dimensions[sr].height is not None:
            ws.row_dimensions[dr].height = ws.row_dimensions[sr].height
    for m in src_merges:
        ws.merge_cells(start_row=m.min_row + spacing, start_column=m.min_col,
                       end_row=m.max_row + spacing, end_column=m.max_col)
    new_date_row = date_row + spacing
    for c in (1, 7, 11):
        ws.cell(new_date_row, c).value = new_date
    changed = apply_inplace(ws, wsv, dst_rows, step, selected_cols, gram_filter)
    rounded = apply_round5(ws, dst_rows, selected_cols)
    wb.save(outfile)
    print("[ok] blok baru %s-%s dibuat (tanggal %s), %d sel diubah, %d dibulatkan ke kelipatan 5"
          % (dst_h, dst_end, new_date, len(changed), rounded))


def do_update(ws, wb, wsv, step, selected_cols, gram_filter, dry_run, outfile):
    headers = find_header_rows(ws)
    if not headers:
        sys.exit("Tidak ada blok 'HARGA ANTAM'.")
    src_h = headers[-1]
    src_end = last_data_row(ws, src_h)
    data_rows = list(range(src_h + 1, src_end + 1))
    print("[info] update di tempat | blok header %d-%d | step %d/gram | target %s" % (src_h, src_end, step, sorted(selected_cols)))
    if dry_run:
        print("[dry-run/update] baris yg akan diubah:")
        for r in data_rows:
            for c in selected_cols:
                g = ws.cell(r, GRAM_COL[c]).value
                try:
                    g = float(g)
                except (TypeError, ValueError):
                    continue
                if gram_filter is not None and g not in gram_filter:
                    continue
                v = ws.cell(r, c).value
                try:
                    base = float(v)
                except (TypeError, ValueError):
                    continue
                print("  row %d | gram %s | col %s | %s -> %s" % (r, g, c, base, base + step * g))
        print("[dry-run] selesai (belum disimpan).")
        return
    changed = apply_inplace(ws, wsv, data_rows, step, selected_cols, gram_filter)
    rounded = apply_round5(ws, data_rows, selected_cols)
    wb.save(outfile)
    print("[ok] blok terakhir diupdate: %d sel diubah, %d dibulatkan" % (len(changed), rounded))


def do_json(ws, wb, json_file, dry_run, outfile):
    with open(json_file, encoding="utf-8") as f:
        data = json.load(f)
    new_date = str(data["date"])
    map_kuning = data.get("kuning", {})
    map_putih = data.get("putih", {})
    label_rows = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str):
            label_rows.setdefault(v.strip(), []).append(r)
    changed = 0

    def put(row, col, val):
        nonlocal changed
        if ws.cell(row, col).value != val:
            ws.cell(row, col).value = val
            changed += 1

    for label, val in map_kuning.items():
        rows = label_rows.get(label.strip())
        if not rows:
            print("[warn] label EMAS tak ditemukan:", label)
            continue
        for rr in rows:
            put(rr, 2, val)
            put(rr, 6, val)
    for label, val in map_putih.items():
        rows = label_rows.get(label.strip())
        if not rows:
            print("[warn] label PUTIH tak ditemukan:", label)
            continue
        for rr in rows:
            put(rr, 4, val)
            put(rr, 8, val)
    for c in (1, 5):
        if ws.cell(1, c).value != new_date:
            ws.cell(1, c).value = new_date
            changed += 1
    if dry_run:
        print("[dry-run/json] %d sel akan diubah, tanggal -> %s" % (changed, new_date))
        return
    wb.save(outfile)
    print("[ok] EMAS diupdate: %d sel, tanggal %s" % (changed, new_date))


def run_cek_anomali(xlsx_path):
    """Cek 2026>2025>RANDOM pada blok terakhir. 0 = bersih, 1 = ada anomali."""
    script = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cek_anomali_antam.py")
    print("[cek] urutan ANTAM 2026>2025>RANDOM (blok terakhir)")
    r = subprocess.run([sys.executable, script, xlsx_path, "--last", "1"])
    return 0 if r.returncode == 0 else 1


def git_commit_push(message):
    try:
        subprocess.run([GIT, "add", "-A"], cwd=REPO, check=True)
        subprocess.run([GIT, "commit", "-m", message], cwd=REPO, check=True)
        subprocess.run([GIT, "push", "origin", "main"], cwd=REPO, check=True)
        print("[git] commit & push selesai.")
    except subprocess.CalledProcessError as e:
        print("[warn] git gagal:", e)


def open_excel(path):
    try:
        os.startfile(path)
    except Exception as e:
        print("[warn] gagal buka Excel:", e)


def send_discord_snapshot(xlsx_path, sheet, caption, channel):
    """Render sheet jadi PNG lalu kirim ke channel Discord (webhook)."""
    tmp = os.path.join(tempfile.gettempdir(), "harga_snap_%s.png" % datetime.now().strftime("%H%M%S"))
    try:
        from screenshot_harga import render
    except ImportError:
        base = os.path.dirname(os.path.abspath(__file__))
        sys.path.insert(0, base)
        from screenshot_harga import render
    try:
        render(xlsx_path, sheet, None, tmp)
    except Exception as e:
        print("[warn] gagal screenshot:", e)
        return
    try:
        subprocess.run(
            [sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)), "send_discord.py"),
             tmp, "--channel", channel, "--caption", caption],
            check=True)
    except subprocess.CalledProcessError as e:
        print("[warn] gagal kirim Discord:", e)
    finally:
        try:
            os.remove(tmp)
        except Exception:
            pass


def main():
    ap = argparse.ArgumentParser(description="Update harga mas (ANTAM + Galeri24 + UBS + EMAS).")
    ap.add_argument("file")
    ap.add_argument("--mode", choices=["add", "update", "json"], default="add")
    ap.add_argument("--date", default=None)
    ap.add_argument("--step", type=int, default=50)
    ap.add_argument("--targets", default=None)
    ap.add_argument("--grams", default=None)
    ap.add_argument("--json", default=None)
    ap.add_argument("--sheet", default="ANTAM")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--no-backup", action="store_true")
    ap.add_argument("--no-open", action="store_true")
    ap.add_argument("--no-git", action="store_true")
    ap.add_argument("--send-discord", action="store_true", help="Kirim foto tabel ke channel Discord setelah update.")
    ap.add_argument("--discord-channel", default=None, help="Nama channel di discord_config.json. Default: harga-lm-emas (LM) / harga-perhiasan (EMAS).")
    args = ap.parse_args()

    if not os.path.exists(args.file):
        sys.exit("File tidak ditemukan: %s" % args.file)
    if args.mode == "add" and not args.date:
        sys.exit("Mode add butuh --date.")
    if args.mode == "json" and not args.json:
        sys.exit("Mode json butuh --json <file>.")

    if not args.no_backup and not args.dry_run:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        base, ext = os.path.splitext(args.file)
        shutil.copy2(args.file, "%s_BACKUP_%s%s" % (base, ts, ext))
        print("[backup] dibuat otomatis.")

    wb = load_workbook(args.file, data_only=False)
    wbv = load_workbook(args.file, data_only=True)
    if args.mode == "json" and args.sheet == "ANTAM":
        args.sheet = "EMAS"
    if args.sheet not in wb.sheetnames:
        sys.exit("Sheet '%s' tidak ada. Ada: %s" % (args.sheet, wb.sheetnames))
    ws = wb[args.sheet]

    labels = set()
    anomali = 0
    if args.mode == "json":
        do_json(ws, wb, args.json, args.dry_run, args.file)
    else:
        selected_cols, labels = resolve_columns(args.targets)
        gram_filter = resolve_grams(args.grams)
        if not selected_cols:
            sys.exit("Tidak ada target valid.")
        if args.mode == "add":
            do_add(ws, wb, wbv, args.step, selected_cols, gram_filter, args.date, args.dry_run, args.file)
        else:
            do_update(ws, wb, wbv, args.step, selected_cols, gram_filter, args.dry_run, args.file)
        if not args.dry_run:
            anomali = run_cek_anomali(args.file)

    if not args.dry_run:
        # Screenshot dulu, sebelum Excel user dibuka (file lock bikin PNG kosong).
        if args.send_discord and anomali:
            print("[skip discord] ada anomali urutan ANTAM — perbaiki dulu.")
        elif args.send_discord:
            sheet = "EMAS" if args.mode == "json" else args.sheet
            if args.discord_channel:
                channel = args.discord_channel
            elif sheet == "EMAS":
                channel = "harga-perhiasan"
            else:
                channel = "harga-lm-emas"
            caption = "Update harga %s - %s" % (sheet, args.date or "EMAS")
            send_discord_snapshot(args.file, sheet, caption, channel)
        if not args.no_open:
            open_excel(args.file)
        if not args.no_git:
            git_commit_push("update harga %s - %s" % (args.date or "EMAS", ", ".join(sorted(labels or ["perhiasan"]))))
    else:
        print("[dry-run] selesai, tidak ada perubahan disimpan.")


if __name__ == "__main__":
    main()

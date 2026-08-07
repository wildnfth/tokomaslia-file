# -*- coding: utf-8 -*-
"""
SKILL: update-harga-mas -> mode EMAS (perhiasan)
================================================
Memperbarui harga perhiasan di sheet EMAS PADA BLOK YANG SUDAH ADA (tidak
menambah blok baru). Mengganti tanggal dan nilai harga kuning & putih sesuai
file JSON, menjaga format (font, size, warna, background, border) tetap sama.
Membuat backup otomatis sebelum menyimpan.

Cara pakai:
  python update_harga_emas.py "<file.xlsx>" --json "<data.json>"

Format JSON:
{
  "date": "6 AGUSTUS 2026",
  "kuning": { "<kadar>": "825 / 850"  atau 1275  ... },
  "putih":  { "<kadar>": 850, ... }
}
  - "kuning"  disalin ke kolom B & F; nilai bisa string "X / Y" atau angka.
  - "putih"   disalin ke kolom D & H; berupa angka (hanya kadar yg ada).
  - kadar = label di kolom A (contoh "300/6K", "9C PABRIK", "9A/916/21K").
  - POT tidak diubah.
"""
import argparse
import json
import os
import shutil
import sys
from datetime import datetime

from openpyxl import load_workbook

# kolom: kuning -> B(2)&F(6); putih -> D(4)&H(8)
KUNING_COLS = (2, 6)
PUTIH_COLS = (4, 8)


def find_row(ws, label, max_row=19):
    for r in range(1, max_row + 1):
        if ws.cell(r, 1).value == label:
            return r
    return None


def apply_values(ws, mapping, cols):
    for kadar, val in mapping.items():
        r = find_row(ws, kadar)
        if r is None:
            print("  [warn] kadar '%s' tidak ditemukan, dilewati." % kadar)
            continue
        for c in cols:
            ws.cell(r, c).value = val


def main():
    ap = argparse.ArgumentParser(description="Update harga perhiasan (sheet EMAS).")
    ap.add_argument("file", help="Path file .xlsx")
    ap.add_argument("--json", required=True, help="Path file JSON berisi nilai baru")
    ap.add_argument("--sheet", default="EMAS", help="Nama sheet (default EMAS)")
    ap.add_argument("--dry-run", action="store_true", help="Cek saja, tanpa menyimpan")
    ap.add_argument("--no-backup", action="store_true", help="Tanpa backup otomatis")
    args = ap.parse_args()

    if not os.path.exists(args.file):
        sys.exit("File tidak ditemukan: %s" % args.file)
    if not os.path.exists(args.json):
        sys.exit("File JSON tidak ditemukan: %s" % args.json)

    with open(args.json, "r", encoding="utf-8") as f:
        data = json.load(f)
    date_new = data.get("date", "")
    kuning = data.get("kuning", {})
    putih = data.get("putih", {})

    if not args.no_backup and not args.dry_run:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        base, ext = os.path.splitext(args.file)
        shutil.copy2(args.file, "%s_BACKUP_%s%s" % (base, ts, ext))
        print("[backup] dibuat otomatis.")

    wb = load_workbook(args.file)
    if args.sheet not in wb.sheetnames:
        sys.exit("Sheet '%s' tidak ditemukan. Ada: %s" % (args.sheet, wb.sheetnames))
    ws = wb[args.sheet]

    if args.dry_run:
        print("[dry-run] sheet %s | tanggal baru: %s" % (args.sheet, date_new or "(tidak diubah)"))
        print("[dry-run] kadar kuning : %d | putih: %d" % (len(kuning), len(putih)))
        return

    if date_new:
        ws["A1"] = date_new + " "
        ws["E1"] = date_new + " "
    apply_values(ws, kuning, KUNING_COLS)
    apply_values(ws, putih, PUTIH_COLS)

    wb.save(args.file)
    print("[ok] %s diupdate (tanggal: %s)" % (args.sheet, date_new or "-"))


if __name__ == "__main__":
    main()

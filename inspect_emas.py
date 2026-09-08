import openpyxl

wb = openpyxl.load_workbook('TEMPLATE HARGA MAS2.xlsx', data_only=True)
ws = wb['EMAS']
print("Max row:", ws.max_row, "Max col:", ws.max_column)
for r in range(1, 18):
    row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    print(f"Row {r:2d}: {row_vals}")
